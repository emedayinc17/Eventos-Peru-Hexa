from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()

# Slide 1: Service IAM Overview
slide_layout = prs.slide_layouts[1]  # Title and Content
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "IAM Service — Resumen"

body = slide.shapes.placeholders[1].text_frame
body.text = "Objetivo"
body.level = 0
p = body.add_paragraph()
p.text = "Autenticación y autorización centralizada para los servicios (login, JWT, gestión de usuarios y roles)."
p.level = 1

p = body.add_paragraph()
p.text = "Características clave:"
p.level = 0
for item in [
    "Emite JWT para consumidores (catalogo, contratacion, proveedores)",
    "Gestión de usuarios: crear/editar/buscar (admin) y cambios de contraseña",
    "Soporta tokens IAM y proveedores externos según configuración",
    "Endpoints REST: /auth, /users, /me, /admin/users, /reset-password"
]:
    q = body.add_paragraph()
    q.text = f"• {item}"
    q.level = 1

# Slide 2: Componentes y Flujo de Autenticación
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "Componentes & Flujo"
body = slide.shapes.placeholders[1].text_frame
body.text = "Componentes principales"
p = body.add_paragraph()
p.text = "• Servicio FastAPI con endpoints REST"
p.level = 1
p = body.add_paragraph()
p.text = "• Módulo compartido 'ev_shared/security.py' para decodificar y validar JWT"
p.level = 1
p = body.add_paragraph()
p.text = "• Base de datos MySQL: esquema ev_iam para usuarios, roles y sesiones"
p.level = 1

p = body.add_paragraph()
p.text = "Flujo de autenticación (resumen):"
p.level = 0
for step in [
    "1) Usuario -> POST /auth/login (email+password)",
    "2) IAM valida credenciales, emite JWT (access + refresh)",
    "3) Consumidores reciben JWT y lo envían en Authorization: Bearer <token>",
    "4) Middleware/cliente valida token usando ev_shared.decode_jwt y aplica roles/permits"
]:
    q = body.add_paragraph()
    q.text = step
    q.level = 1

# Save
prs.save('IAM_Service_Resumen.pptx')
print('PPTX generado: IAM_Service_Resumen.pptx')
